"""PresenterAgent — présentations niveau Gamma — Phase 10.

Produit DEUX livrables à partir d'un sujet (ou d'un texte fourni) :

    1. Un deck HTML animé et thémé (le cœur de l'expérience, façon Gamma) :
       navigation clavier/clic, transitions, puces révélées en cascade,
       images IA contextuelles par diapositive (Pollinations — générées
       par le NAVIGATEUR au chargement : zéro réseau côté serveur).
       → sert d'APERÇU instantané dans l'interface avant tout téléchargement.
    2. Un fichier .pptx thémé (export PowerPoint classique, avec notes).

Le LLM produit un plan JSON unique contenant, pour chaque diapositive,
titre + puces + notes + image_prompt (description anglaise pour l'illustration).
Repli déterministe si tous les providers échouent : les fichiers sortent
quand même.

Contrat BaseAgent respecté (jamais d'exception, duration_ms renseigné).
La logique de génération est exposée en fonction module-level
(generate_presentation) pour être appelée aussi par le router HTTP du
wizard (/presentations/generate) sans passer par l'orchestrateur.
"""
from __future__ import annotations

import html as html_lib
import json
import re
import time
import urllib.parse
import uuid
from pathlib import Path
from typing import Any

from core.agents.base import AgentAutonomy, BaseAgent
from core.agents.models import AgentResult, AgentTask, ExecutionContext
from core.provider_layer.router import ALL_PROVIDERS_FAILED, get_router

OUTPUTS_DIR = Path("brain_data/outputs")

# ── Thèmes (partagés HTML + PPTX) ─────────────────────────────────────────────
THEMES: dict[str, dict[str, Any]] = {
    "makenbrain": {"label": "Sombre turquoise", "bg": "#0b0e14", "bg2": "#11151d",
                   "text": "#e8edf2", "muted": "#9aa3ad", "accent": "#2de2c9",
                   "pptx_bg": (11, 14, 20), "pptx_text": (232, 237, 242),
                   "pptx_muted": (154, 163, 173), "pptx_accent": (45, 226, 201)},
    "clair":      {"label": "Clair professionnel", "bg": "#f7f8fa", "bg2": "#ffffff",
                   "text": "#111827", "muted": "#4b5563", "accent": "#2563eb",
                   "pptx_bg": (247, 248, 250), "pptx_text": (17, 24, 39),
                   "pptx_muted": (75, 85, 99), "pptx_accent": (37, 99, 235)},
    "sable":      {"label": "Sable chaleureux", "bg": "#f5efe6", "bg2": "#fbf7f0",
                   "text": "#292018", "muted": "#6b5d4f", "accent": "#c2410c",
                   "pptx_bg": (245, 239, 230), "pptx_text": (41, 32, 24),
                   "pptx_muted": (107, 93, 79), "pptx_accent": (194, 65, 12)},
    "violet":     {"label": "Nuit violette", "bg": "#120a1f", "bg2": "#1a1029",
                   "text": "#ece8f5", "muted": "#a89fc0", "accent": "#a78bfa",
                   "pptx_bg": (18, 10, 31), "pptx_text": (236, 232, 245),
                   "pptx_muted": (168, 159, 192), "pptx_accent": (167, 139, 250)},
}

FONTS: dict[str, str] = {
    "moderne":  "'Segoe UI', system-ui, -apple-system, sans-serif",
    "elegante": "Georgia, 'Times New Roman', serif",
    "mono":     "'Cascadia Code', 'JetBrains Mono', Consolas, monospace",
}

_OUTLINE_SYSTEM_PROMPT = """Tu es un concepteur de présentations professionnelles de très haut niveau.

Tu produis la STRUCTURE d'une présentation percutante, réaliste et documentée — jamais de contenu inventé présenté comme un fait précis (pas de faux chiffres, pas de fausses citations).

Réponds UNIQUEMENT avec un objet JSON valide, sans texte avant ni après :
{
  "title": "Titre de la présentation",
  "subtitle": "Sous-titre en une phrase",
  "slides": [
    {
      "title": "Titre de la diapositive",
      "bullets": ["Point percutant (12 mots max)", "..."],
      "notes": "Notes du présentateur : ce qu'il faut DIRE (3-4 phrases riches).",
      "image_prompt": "professional photography of ..., short English description to illustrate this slide"
    }
  ]
}

Règles :
- Structure narrative : accroche, contexte/problème, points clés développés, exemples concrets, conclusion avec appel à l'action.
- 3 à 5 puces par diapositive, courtes — le détail vit dans les notes.
- image_prompt : EN ANGLAIS, descriptif, photographique et directement lié au contenu de la diapo (jamais générique).
- Si un TEXTE SOURCE est fourni, la présentation le suit FIDÈLEMENT — n'invente rien qui le contredise.
- Respecte STRICTEMENT la langue et le nombre de diapositives demandés."""


# ══════════════════════════════════════════════════════════════════════════════
# Génération (module-level — utilisée par l'agent ET le router du wizard)
# ══════════════════════════════════════════════════════════════════════════════

async def generate_presentation(
    topic: str,
    slides_count: int | None = None,
    theme: str = "makenbrain",
    font: str = "moderne",
    with_images: bool = True,
    source_text: str | None = None,
    language: str | None = None,
) -> dict[str, Any]:
    """Génère le deck HTML animé + l'export .pptx et retourne leurs URLs.

    Returns:
        {"title", "subtitle", "slides_count", "html_url", "pptx_url", "strategy"}
    """
    theme_cfg = THEMES.get(theme, THEMES["makenbrain"])
    font_css = FONTS.get(font, FONTS["moderne"])

    outline, strategy = await _build_outline(topic, slides_count, source_text, language)

    base = _safe_name(outline.get("title", "presentation"))
    token = uuid.uuid4().hex[:8]
    html_name = f"{base}-{token}.html"
    pptx_name = f"{base}-{token}.pptx"

    OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)
    (OUTPUTS_DIR / html_name).write_text(
        _build_html_deck(outline, theme_cfg, font_css, with_images), encoding="utf-8"
    )
    _write_pptx(outline, theme_cfg, OUTPUTS_DIR / pptx_name)

    return {
        "title": outline.get("title", "Présentation"),
        "subtitle": outline.get("subtitle", ""),
        "slides_count": len(outline.get("slides", [])),
        "slide_titles": [s.get("title", "") for s in outline.get("slides", [])],
        "html_url": f"/static/outputs/{html_name}",
        "pptx_url": f"/static/outputs/{pptx_name}",
        "strategy": strategy,
    }


async def _build_outline(
    topic: str,
    slides_count: int | None,
    source_text: str | None,
    language: str | None,
) -> tuple[dict[str, Any], str]:
    """Structure via LLM ; repli déterministe si tous les providers échouent."""
    count = max(3, min(slides_count or 8, 15))
    parts = [f"Sujet de la présentation :\n{topic}"]
    parts.append(f"Nombre de diapositives demandé (hors titre) : {count}.")
    parts.append(f"Langue : {language or 'celle du sujet (français par défaut)'}.")
    if source_text and source_text.strip():
        parts.append(f"TEXTE SOURCE à suivre fidèlement :\n{source_text.strip()[:6000]}")
    try:
        raw = await get_router().generate(
            prompt="\n\n".join(parts),
            system_prompt=_OUTLINE_SYSTEM_PROMPT,
        )
        if raw and raw != ALL_PROVIDERS_FAILED:
            match = re.search(r"\{.*\}", raw, re.DOTALL)
            if match:
                outline = json.loads(match.group())
                if outline.get("slides"):
                    outline["slides"] = outline["slides"][:15]
                    return outline, "llm"
    except Exception:
        pass
    return _fallback_outline(topic, count), "fallback"


def _fallback_outline(topic: str, count: int) -> dict[str, Any]:
    """Plan minimal si aucun LLM — le fichier sort quand même."""
    titles = ["Introduction", "Contexte", "Points clés", "Analyse",
              "Exemples", "Opportunités", "Risques", "Conclusion"][:max(3, min(count, 8))]
    return {
        "title": topic.strip()[:120] or "Présentation",
        "subtitle": "Structure générée hors-ligne — à compléter",
        "slides": [
            {"title": t, "bullets": ["Point 1", "Point 2", "Point 3"],
             "notes": "", "image_prompt": f"professional photo, {topic[:60]}"}
            for t in titles
        ],
    }


def _safe_name(title: str) -> str:
    return re.sub(r"[^a-zA-Z0-9_-]+", "-", str(title))[:40].strip("-") or "presentation"


def _pollinations_url(prompt: str, seed: int) -> str:
    encoded = urllib.parse.quote((prompt or "professional abstract background")[:300])
    return (f"https://image.pollinations.ai/prompt/{encoded}"
            f"?width=896&height=640&nologo=true&enhance=true&seed={seed}")


# ── Deck HTML animé (autonome, aucun CDN requis) ──────────────────────────────

def _build_html_deck(
    outline: dict[str, Any],
    theme: dict[str, Any],
    font_css: str,
    with_images: bool,
) -> str:
    """Construit le deck HTML : transitions, puces en cascade, images IA, clavier."""
    esc = html_lib.escape
    slides_html: list[str] = []

    # Diapositive de titre
    slides_html.append(
        '<section class="slide title-slide">'
        f'<h1>{esc(str(outline.get("title", "Présentation"))[:140])}</h1>'
        f'<p class="subtitle">{esc(str(outline.get("subtitle", ""))[:200])}</p>'
        '<p class="brand">Généré par MakenBrain</p>'
        "</section>"
    )

    for idx, s in enumerate(outline.get("slides", [])[:15]):
        bullets = "".join(
            f'<li style="animation-delay:{0.15 * (i + 1):.2f}s">{esc(str(b)[:200])}</li>'
            for i, b in enumerate([b for b in s.get("bullets", []) if str(b).strip()][:6])
        )
        img_html = ""
        layout = ""
        if with_images and s.get("image_prompt"):
            url = _pollinations_url(str(s["image_prompt"]), seed=1000 + idx)
            img_html = (
                '<figure class="slide-img">'
                f'<img src="{esc(url)}" alt="" loading="lazy" onerror="this.closest(\'figure\').style.display=\'none\'">'
                "</figure>"
            )
            layout = " with-img"
        slides_html.append(
            f'<section class="slide{layout}">'
            f"<h2>{esc(str(s.get('title', ''))[:120])}</h2>"
            f'<div class="slide-body"><ul>{bullets}</ul>{img_html}</div>'
            "</section>"
        )

    total = len(slides_html)
    dots = "".join(f'<span class="dot" data-i="{i}"></span>' for i in range(total))

    return f"""<!DOCTYPE html>
<html lang="fr">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{esc(str(outline.get("title", "Présentation")))}</title>
<style>
  :root {{
    --bg: {theme['bg']}; --bg2: {theme['bg2']}; --text: {theme['text']};
    --muted: {theme['muted']}; --accent: {theme['accent']};
  }}
  * {{ margin:0; padding:0; box-sizing:border-box; }}
  html, body {{ height:100%; overflow:hidden; background:var(--bg);
    color:var(--text); font-family:{font_css}; }}
  .deck {{ position:relative; height:100%; }}
  .slide {{
    position:absolute; inset:0; padding:7vh 9vw; display:flex; flex-direction:column;
    justify-content:center; opacity:0; visibility:hidden; transform:translateY(26px) scale(.985);
    transition:opacity .55s ease, transform .55s ease, visibility .55s;
    background:linear-gradient(150deg, var(--bg) 0%, var(--bg2) 100%);
  }}
  .slide.active {{ opacity:1; visibility:visible; transform:none; }}
  .slide::before {{ content:''; position:absolute; left:0; top:0; bottom:0; width:7px;
    background:var(--accent); }}
  .title-slide {{ align-items:flex-start; }}
  h1 {{ font-size:clamp(2.2rem, 5.5vw, 4.2rem); line-height:1.12; max-width:22ch; }}
  .subtitle {{ margin-top:1.2rem; font-size:clamp(1rem, 2vw, 1.5rem); color:var(--accent); }}
  .brand {{ position:absolute; bottom:5vh; font-size:.8rem; color:var(--muted);
    letter-spacing:2px; text-transform:uppercase; }}
  h2 {{ font-size:clamp(1.6rem, 3.4vw, 2.6rem); color:var(--accent); margin-bottom:4vh; }}
  .slide-body {{ display:flex; gap:5vw; align-items:center; }}
  .slide-body ul {{ list-style:none; flex:1; }}
  .slide-body li {{
    font-size:clamp(1rem, 1.9vw, 1.45rem); line-height:1.5; margin-bottom:2.2vh;
    padding-left:1.6em; position:relative; opacity:0; transform:translateX(-14px);
  }}
  .slide.active li {{ animation:reveal .5s ease forwards; }}
  @keyframes reveal {{ to {{ opacity:1; transform:none; }} }}
  .slide-body li::before {{ content:'▸'; position:absolute; left:0; color:var(--accent); }}
  .slide-img {{ flex:0 0 34%; max-width:34%; }}
  .slide-img img {{ width:100%; border-radius:18px; box-shadow:0 24px 60px rgba(0,0,0,.35);
    aspect-ratio:7/5; object-fit:cover; background:var(--bg2); }}
  .nav {{ position:fixed; bottom:3vh; left:50%; transform:translateX(-50%);
    display:flex; gap:9px; z-index:10; }}
  .dot {{ width:9px; height:9px; border-radius:50%; background:var(--muted);
    opacity:.35; cursor:pointer; transition:.25s; }}
  .dot.on {{ background:var(--accent); opacity:1; transform:scale(1.3); }}
  .counter {{ position:fixed; bottom:2.6vh; right:3vw; font-size:.85rem; color:var(--muted); }}
  .arrow {{ position:fixed; top:50%; transform:translateY(-50%); z-index:10; border:none;
    background:transparent; color:var(--muted); font-size:2.2rem; cursor:pointer;
    padding:1rem; transition:.2s; }}
  .arrow:hover {{ color:var(--accent); }}
  #prev {{ left:1vw; }} #next {{ right:1vw; }}
  .progress {{ position:fixed; top:0; left:0; height:3px; background:var(--accent);
    transition:width .4s ease; z-index:10; }}
</style>
</head>
<body>
<div class="progress" id="progress"></div>
<div class="deck">{''.join(slides_html)}</div>
<button class="arrow" id="prev" aria-label="Précédent">‹</button>
<button class="arrow" id="next" aria-label="Suivant">›</button>
<div class="nav">{dots}</div>
<div class="counter"><span id="cur">1</span> / {total}</div>
<script>
  const slides = document.querySelectorAll('.slide');
  const dots = document.querySelectorAll('.dot');
  let i = 0;
  function show(n) {{
    i = Math.max(0, Math.min(n, slides.length - 1));
    slides.forEach((s, k) => s.classList.toggle('active', k === i));
    dots.forEach((d, k) => d.classList.toggle('on', k === i));
    document.getElementById('cur').textContent = i + 1;
    document.getElementById('progress').style.width = ((i + 1) / slides.length * 100) + '%';
  }}
  document.getElementById('next').onclick = () => show(i + 1);
  document.getElementById('prev').onclick = () => show(i - 1);
  dots.forEach(d => d.onclick = () => show(+d.dataset.i));
  document.addEventListener('keydown', e => {{
    if (e.key === 'ArrowRight' || e.key === ' ' || e.key === 'PageDown') show(i + 1);
    if (e.key === 'ArrowLeft' || e.key === 'PageUp') show(i - 1);
    if (e.key === 'Home') show(0);
    if (e.key === 'End') show(slides.length - 1);
  }});
  show(0);
</script>
</body>
</html>"""


# ── Export .pptx thémé ─────────────────────────────────────────────────────────

def _write_pptx(outline: dict[str, Any], theme: dict[str, Any], path: Path) -> None:
    """Construit le .pptx aux couleurs du thème choisi (avec notes)."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    bg, text_c = RGBColor(*theme["pptx_bg"]), RGBColor(*theme["pptx_text"])
    muted_c, accent_c = RGBColor(*theme["pptx_muted"]), RGBColor(*theme["pptx_accent"])

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def decorate(slide) -> None:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = bg
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                                     Inches(0.12), prs.slide_height)
        bar.fill.solid(); bar.fill.fore_color.rgb = accent_c; bar.line.fill.background()

    slide = prs.slides.add_slide(blank)
    decorate(slide)
    box = slide.shapes.add_textbox(Inches(0.9), Inches(2.6), Inches(11.5), Inches(1.8))
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = str(outline.get("title", "Présentation"))[:120]
    p.font.size = Pt(44); p.font.bold = True; p.font.color.rgb = text_c
    sub = tf.add_paragraph()
    sub.text = str(outline.get("subtitle", ""))[:160]
    sub.font.size = Pt(20); sub.font.color.rgb = accent_c
    footer = slide.shapes.add_textbox(Inches(0.9), Inches(6.8), Inches(6), Inches(0.4))
    fp = footer.text_frame.paragraphs[0]
    fp.text = "Généré par MakenBrain"; fp.font.size = Pt(11); fp.font.color.rgb = muted_c

    for s in outline.get("slides", [])[:15]:
        slide = prs.slides.add_slide(blank)
        decorate(slide)
        tp = slide.shapes.add_textbox(Inches(0.9), Inches(0.5), Inches(11.5), Inches(1.0)) \
                         .text_frame.paragraphs[0]
        tp.text = str(s.get("title", ""))[:110]
        tp.font.size = Pt(30); tp.font.bold = True; tp.font.color.rgb = accent_c
        btf = slide.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(11.0), Inches(5.0)).text_frame
        btf.word_wrap = True
        bullets = [str(b) for b in s.get("bullets", []) if str(b).strip()][:6]
        for i, bullet in enumerate(bullets):
            para = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
            para.text = f"•  {bullet[:180]}"
            para.font.size = Pt(20); para.font.color.rgb = text_c
            para.space_after = Pt(14)
        notes = str(s.get("notes", "")).strip()
        if notes:
            slide.notes_slide.notes_text_frame.text = notes[:2000]

    prs.save(str(path))


# ══════════════════════════════════════════════════════════════════════════════
# Agent (Salle des Agents + orchestrateur)
# ══════════════════════════════════════════════════════════════════════════════

class PresenterAgent(BaseAgent):
    """Agent présentateur — deck HTML animé (aperçu) + export .pptx."""

    name        = "presenter_agent"
    description = (
        "Concepteur de présentations niveau pro : deck HTML animé avec images "
        "IA contextuelles (aperçu instantané) + export PowerPoint (.pptx) "
        "avec notes du présentateur."
    )
    capabilities = ["presentation", "powerpoint", "pptx", "slides", "diapo"]
    autonomy     = AgentAutonomy.SANDBOXED_EXECUTE
    version      = "2.0.0"

    cost_per_call        = 1.0
    confidence_threshold = 0.75

    async def run(self, task: AgentTask, ctx: ExecutionContext) -> AgentResult:
        t0 = time.monotonic()
        try:
            result = await generate_presentation(
                topic       = task.input,
                slides_count= task.context.get("slides_count"),
                theme       = task.context.get("theme", "makenbrain"),
                font        = task.context.get("font", "moderne"),
                with_images = bool(task.context.get("with_images", True)),
                source_text = task.context.get("source_text"),
                language    = task.context.get("language"),
            )
            titles = "\n".join(f"{i+1}. {t}" for i, t in enumerate(result["slide_titles"]))
            output = (
                f"## 📊 {result['title']}\n\n*{result['subtitle']}*\n\n"
                f"**{result['slides_count']} diapositives** :\n\n{titles}\n\n"
                f"**[▶ Aperçu animé (HTML)]({result['html_url']})** · "
                f"**[⬇ PowerPoint (.pptx)]({result['pptx_url']})**"
            )
            ctx.shared[f"{self.name}_output"] = output
            return self._timed_result(
                task, t0, success=True, output=output,
                confidence=0.85 if result["strategy"] == "llm" else 0.55,
                metadata=result,
            )
        except Exception as exc:  # noqa: BLE001 — contrat BaseAgent : ne jamais lever.
            return self._timed_result(task, t0, success=False, error=str(exc))
